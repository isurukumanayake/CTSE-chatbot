import os
import time  # Added missing import
from pathlib import Path
from typing import List, Optional
import tempfile

from langchain.docstore.document import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, UnstructuredPowerPointLoader
from langchain_community.document_loaders.word_document import UnstructuredWordDocumentLoader

from ..utils.logger import get_logger
from config import Config

class DocumentProcessor:
    """Handle document loading and processing"""
    
    def __init__(self):
        self.logger = get_logger(__name__)
        self.config = Config()
        
        # Initialize text splitter
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.config.CHUNK_SIZE,
            chunk_overlap=self.config.CHUNK_OVERLAP,
            length_function=len,
            separators=["\n\n", "\n", ".", " ", ""]
        )
    
    def process_file(self, file_path: Path) -> List[Document]:
        """Process a single file and return document chunks"""
        try:
            self.logger.info(f"Processing file: {file_path.name}")
            
            # Validate file
            if not self._validate_file(file_path):
                return []
            
            # Load document based on file type
            documents = self._load_document(file_path)
            
            if not documents:
                self.logger.warning(f"No content loaded from {file_path.name}")
                return []
            
            # Split into chunks
            chunks = self.text_splitter.split_documents(documents)
            
            # Add metadata
            for chunk in chunks:
                chunk.metadata.update({
                    "source": str(file_path.name),
                    "file_type": file_path.suffix.lower(),
                    "processed_at": time.time()
                })
            
            self.logger.info(f"Created {len(chunks)} chunks from {file_path.name}")
            return chunks
            
        except Exception as e:
            self.logger.error(f"Error processing {file_path.name}: {e}")
            return []
    
    def _validate_file(self, file_path: Path) -> bool:
        """Validate file before processing"""
        # Check if file exists
        if not file_path.exists():
            self.logger.error(f"File not found: {file_path}")
            return False
        
        # Check file extension
        if file_path.suffix.lower() not in self.config.ALLOWED_EXTENSIONS:
            self.logger.error(f"Unsupported file type: {file_path.suffix}")
            return False
        
        # Check file size
        if file_path.stat().st_size > self.config.MAX_FILE_SIZE:
            self.logger.error(f"File too large: {file_path.name}")
            return False
        
        return True
    
    def _load_document(self, file_path: Path) -> List[Document]:
        """Load document based on file type"""
        file_extension = file_path.suffix.lower()
        
        try:
            if file_extension == '.pdf':
                return self._load_pdf(file_path)
            elif file_extension == '.pptx':
                return self._load_pptx(file_path)
            elif file_extension == '.docx':
                return self._load_docx(file_path)
            elif file_extension == '.txt':
                return self._load_txt(file_path)
            else:
                self.logger.error(f"Unsupported file type: {file_extension}")
                return []
                
        except Exception as e:
            self.logger.error(f"Failed to load {file_path.name}: {e}")
            return []
    
    def _load_pdf(self, file_path: Path) -> List[Document]:
        """Load PDF document"""
        try:
            loader = PyPDFLoader(str(file_path))
            documents = loader.load()
            self.logger.info(f"Loaded {len(documents)} pages from PDF: {file_path.name}")
            return documents
        except Exception as e:
            self.logger.error(f"PDF loading error for {file_path.name}: {e}")
            return []
    
    def _load_pptx(self, file_path: Path) -> List[Document]:
        """Load PowerPoint document"""
        try:
            loader = UnstructuredPowerPointLoader(str(file_path))
            documents = loader.load()
            self.logger.info(f"Loaded PowerPoint: {file_path.name}")
            return documents
        except Exception as e:
            self.logger.error(f"PowerPoint loading error for {file_path.name}: {e}")
            # Try alternative method for PPTX
            return self._load_pptx_alternative(file_path)
    
    def _load_pptx_alternative(self, file_path: Path) -> List[Document]:
        """Alternative PPTX loader using python-pptx directly"""
        try:
            from pptx import Presentation
            
            prs = Presentation(str(file_path))
            documents = []
            
            for slide_num, slide in enumerate(prs.slides, 1):
                # Extract text from slide
                slide_text = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text += shape.text + "\n"
                
                if slide_text.strip():
                    doc = Document(
                        page_content=slide_text.strip(),
                        metadata={
                            "source": str(file_path.name),
                            "slide": slide_num,
                            "total_slides": len(prs.slides)
                        }
                    )
                    documents.append(doc)
            
            self.logger.info(f"Loaded {len(documents)} slides from PowerPoint: {file_path.name}")
            return documents
            
        except Exception as e:
            self.logger.error(f"Alternative PowerPoint loading failed for {file_path.name}: {e}")
            return []
    
    def _load_docx(self, file_path: Path) -> List[Document]:
        """Load Word document"""
        try:
            loader = UnstructuredWordDocumentLoader(str(file_path))
            documents = loader.load()
            self.logger.info(f"Loaded Word document: {file_path.name}")
            return documents
        except Exception as e:
            self.logger.error(f"Word document loading error for {file_path.name}: {e}")
            return []
    
    def _load_txt(self, file_path: Path) -> List[Document]:
        """Load text document"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            if content.strip():
                doc = Document(
                    page_content=content,
                    metadata={"source": str(file_path.name)}
                )
                self.logger.info(f"Loaded text file: {file_path.name}")
                return [doc]
            else:
                self.logger.warning(f"Empty text file: {file_path.name}")
                return []
                
        except Exception as e:
            self.logger.error(f"Text file loading error for {file_path.name}: {e}")
            return []